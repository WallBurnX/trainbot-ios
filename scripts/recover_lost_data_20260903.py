#!/usr/bin/env python3
"""Restore the rows that disappeared from the live database on 24.07.2026.

What happened: the production SQLite file was replaced by a 14.07 snapshot
during the training-plan release, so every row logged between 15.07 and
23.07 survives only in the Excel export the bot produced on 24.07 00:22
(``Leonid_Fitness_Tracker (2).xlsx``).  Two later workouts (31.08 and 03.09)
were typed into Telegram while the bot was stuck waiting for a date, so they
exist only in ``chat_history``.

The script is idempotent: every row is checked against the target database
before insertion and nothing is deleted.  Dry-run is the default.

    # inside the container (scripts/ is copied into the image):
    python /app/scripts/recover_lost_data_20260903.py \
        --db /app/data/trainbot.db \
        --excel /app/files/recovery/Leonid_Fitness_Tracker_20260724.xlsx --apply
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import datetime

BACKEND_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "backend")
sys.path.insert(0, BACKEND_DIR)

CUTOFF_DATE = "2026-07-14"  # rows on or before this date already exist


# Workouts typed into Telegram after the data loss.  Names follow the labels
# the CRM already uses for the same machines; the raw wording is preserved.
CHAT_WORKOUTS = [
    {
        "date": "2026-08-31", "time": "09:00",
        "exercise": "Жим гантелей", "exercise_raw": "Жим гантелей",
        "weight": "20/24/26/26/26", "sets": 5, "reps": "10/8/8/8/7",
        "comment": "Грудь. Восстановлено из сообщения Telegram 31.08 09:21.",
    },
    {
        "date": "2026-08-31", "time": "09:00",
        "exercise": "Жим в тренажёре Chest Press", "exercise_raw": "Chest press",
        "weight": "57.5/65/70/60", "sets": 4, "reps": "9/7/4/5",
        "comment": "Восстановлено из сообщения Telegram 31.08 09:21.",
    },
    {
        "date": "2026-09-03", "time": "19:00",
        "exercise": "Тяга верхнего блока", "exercise_raw": "Пулл даун",
        "weight": "87.5", "sets": 3, "reps": "10/8/9",
        "comment": "После работы. Восстановлено из сообщения Telegram 03.09 20:15.",
    },
    {
        "date": "2026-09-03", "time": "19:00",
        "exercise": "Горизонтальная тяга на верх спины в тренажёре", "exercise_raw": "Upperback",
        "weight": "30/40/50/60", "sets": 4, "reps": "12/10/10/7",
        "comment": "Восстановлено из сообщения Telegram 03.09 20:15.",
    },
    {
        "date": "2026-09-03", "time": "19:00",
        "exercise": "Тяга горизонтального блока сидя", "exercise_raw": "Лоу роу",
        "weight": "75/85/85", "sets": 3, "reps": "6+5/4+4/3+4",
        "comment": "Low Row: после подхода на обе руки сразу 5/4/4 повторений на каждую руку. "
                   "Восстановлено из сообщения Telegram 03.09 20:15.",
    },
    {
        "date": "2026-09-03", "time": "19:00",
        "exercise": "Разведение рук назад в тренажёре", "exercise_raw": "Обратная бабочка",
        "weight": "15", "sets": 1, "reps": "10-12",
        "comment": "Не пошло, плечи устали. Восстановлено из сообщения Telegram 03.09 20:15.",
    },
    {
        "date": "2026-09-03", "time": "19:00",
        "exercise": "Жим в тренажёре Chest Press", "exercise_raw": "Chest press",
        "weight": "50/70/80", "sets": 3, "reps": "12/10/5",
        "comment": "Восстановлено из сообщения Telegram 03.09 20:15.",
    },
    {
        "date": "2026-09-03", "time": "19:00",
        "exercise": "Сведение рук в тренажёре", "exercise_raw": "Пекторал",
        "weight": "25/30/35", "sets": 3, "reps": "12/10/8",
        "comment": "Восстановлено из сообщения Telegram 03.09 20:15.",
    },
    {
        "date": "2026-09-03", "time": "19:00",
        "exercise": "Гиперэкстензия", "exercise_raw": "Экстензия",
        "weight": "15", "sets": 3, "reps": None,
        "comment": "«Экстензия 15 х 3»: 15 кг, 3 подхода, повторы не указаны. "
                   "Восстановлено из сообщения Telegram 03.09 20:15.",
    },
]


def _iso(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _norm(value) -> str:
    return str(value or "").strip().replace(",", ".").lower()


def _tonnage(weight, sets, reps, fallback):
    try:
        import ai  # requires the provider SDKs; available inside the container
        calculated = float(ai.calculate_tonnage(weight, sets, reps) or 0)
        if calculated > 0:
            return calculated
    except Exception:  # pragma: no cover - local runs without SDKs
        pass
    try:
        return float(fallback or 0)
    except (TypeError, ValueError):
        return 0.0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--db", required=True, help="Path to trainbot.db")
    parser.add_argument("--excel", required=True, help="Excel export produced on 24.07.2026")
    parser.add_argument("--apply", action="store_true", help="Write changes (default: dry-run)")
    args = parser.parse_args()

    os.environ["DB_PATH"] = os.path.abspath(args.db)
    import openpyxl
    import database as db
    from exercises import exercise_alias_key

    db.DB_PATH = os.path.abspath(args.db)
    integrity = db.get_db_connection().execute("PRAGMA integrity_check").fetchone()[0]
    if integrity != "ok":
        raise SystemExit(f"Database integrity check failed: {integrity}")
    db.init_db()

    workbook = openpyxl.load_workbook(args.excel, data_only=True)
    summary = {"workouts": 0, "nutrition": 0, "measurements": 0, "wellness": 0, "schedule": 0, "skipped": 0}

    def workout_exists(conn, row: dict) -> bool:
        keys = {exercise_alias_key(row.get("exercise") or ""), exercise_alias_key(row.get("exercise_raw") or "")}
        keys.discard("")
        for existing in conn.execute(
            "SELECT exercise, exercise_raw, exercise_key, weight, reps FROM workouts WHERE date = ?",
            (row["date"],),
        ).fetchall():
            existing_keys = {
                existing["exercise_key"] or "",
                exercise_alias_key(existing["exercise"] or ""),
                exercise_alias_key(existing["exercise_raw"] or ""),
            }
            existing_keys.discard("")
            if (
                keys & existing_keys
                and _norm(existing["weight"]) == _norm(row.get("weight"))
                and _norm(existing["reps"]) == _norm(row.get("reps"))
            ):
                return True
        return False

    pending_workouts: list[dict] = []
    sheet = workbook["Тренировки"]
    header = [str(cell or "").strip() for cell in next(sheet.iter_rows(values_only=True))]
    index = {name: position for position, name in enumerate(header)}
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        date = _iso(raw[index["Дата"]])
        if not date or date <= CUTOFF_DATE:
            continue
        pending_workouts.append(
            {
                "date": date,
                "time": _text(raw[index["Время"]]),
                "exercise": _text(raw[index["Упражнение"]]),
                "exercise_raw": _text(raw[index.get("Исходное название", index["Упражнение"])]),
                "weight": _text(raw[index["Вес"]]),
                "sets": raw[index["Подходы"]],
                "reps": _text(raw[index["Повторения"]]),
                "tonnage": raw[index["Тоннаж"]],
                "treadmill": _text(raw[index["Дорожка (скор/наклон)"]]),
                "pulse": _text(raw[index["Пульс макс/средний"]]),
                "comment": _text(raw[index["Комментарий"]]),
            }
        )
    pending_workouts.extend(dict(item) for item in CHAT_WORKOUTS)

    with db.db_connection() as conn:
        for row in pending_workouts:
            if not row.get("exercise"):
                continue
            if workout_exists(conn, row):
                summary["skipped"] += 1
                continue
            try:
                row["sets"] = int(float(row["sets"])) if row.get("sets") not in (None, "") else None
            except (TypeError, ValueError):
                row["sets"] = None
            row["tonnage"] = _tonnage(row.get("weight"), row.get("sets"), row.get("reps"), row.get("tonnage"))
            print(f"[workout] {row['date']} {row['exercise']} {row.get('weight') or '-'} × {row.get('reps') or '-'} → {row['tonnage']:.0f} кг")
            summary["workouts"] += 1
            if args.apply:
                db.save_workout(row)

    sheet = workbook["Питание"]
    with db.db_connection() as conn:
        for raw in sheet.iter_rows(min_row=2, values_only=True):
            date = _iso(raw[0])
            if not date or date <= CUTOFF_DATE:
                continue
            meal = _text(raw[2])
            if not meal:
                continue
            calories = raw[6]
            exists = conn.execute(
                """
                SELECT 1 FROM nutrition
                WHERE date = ? AND LOWER(meal) = LOWER(?) AND COALESCE(calories, -1) = COALESCE(?, -1)
                LIMIT 1
                """,
                (date, meal, calories),
            ).fetchone()
            if exists:
                summary["skipped"] += 1
                continue
            print(f"[nutrition] {date} {meal[:60]} · {raw[3]}/{raw[4]}/{raw[5]} · {calories} ккал")
            summary["nutrition"] += 1
            if args.apply:
                db.save_nutrition(
                    {
                        "date": date, "time": _text(raw[1]), "meal": meal,
                        "proteins": raw[3], "fats": raw[4], "carbs": raw[5],
                        "calories": calories, "notes": _text(raw[7]),
                    }
                )

    sheet = workbook["Замеры"]
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        date = _iso(raw[0])
        if not date or date <= "2026-07-12":
            continue
        payload = {
            "date": date, "weight": raw[1], "chest": raw[2], "waist": raw[3], "abdomen": raw[4],
            "thigh": raw[5], "arm": raw[6], "smm": raw[7], "fat_kg": raw[8], "fat_pct": raw[9],
            "visceral": raw[10],
        }
        existing = db.get_measurements_by_date(date)
        if existing and all(
            existing[0].get(key) is not None or value is None for key, value in payload.items() if key != "date"
        ):
            summary["skipped"] += 1
            continue
        print(f"[measurement] {date} {', '.join(f'{k}={v}' for k, v in payload.items() if k != 'date' and v is not None)}")
        summary["measurements"] += 1
        if args.apply:
            db.save_measurement(payload)

    sheet = workbook["Самочувствие"]
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        date = _iso(raw[0])
        if not date or date <= CUTOFF_DATE:
            continue
        payload = {
            "date": date, "sleep_hours": raw[1], "sleep_quality": raw[2], "steps": raw[3],
            "water_l": raw[4], "mood": raw[5], "energy": raw[6], "resting_pulse": raw[7],
            "notes": _text(raw[8]),
        }
        existing = [item for item in db.get_wellness(limit=400) if item.get("date") == date]
        if existing:
            summary["skipped"] += 1
            continue
        print(f"[wellness] {date} sleep={raw[1]} notes={str(raw[8] or '')[:60]}")
        summary["wellness"] += 1
        if args.apply:
            db.save_wellness(payload)

    sheet = workbook["Расписание"]
    with db.db_connection() as conn:
        for raw in sheet.iter_rows(min_row=2, values_only=True):
            date = _iso(raw[0])
            title = _text(raw[2])
            if not date or not title or date <= "2026-07-10":
                continue
            exists = conn.execute(
                "SELECT 1 FROM schedule WHERE date = ? AND title = ? LIMIT 1", (date, title)
            ).fetchone()
            if exists:
                summary["skipped"] += 1
                continue
            print(f"[schedule] {date} {title}")
            summary["schedule"] += 1
            if args.apply:
                db.save_schedule_item(
                    {
                        "date": date, "time": _text(raw[1]), "title": title, "category": _text(raw[3]),
                        "status": _text(raw[4]) or "planned", "notes": _text(raw[5]),
                    }
                )

    mode = "APPLIED" if args.apply else "DRY-RUN"
    print(f"\n{mode}: " + ", ".join(f"{key}={value}" for key, value in summary.items()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
